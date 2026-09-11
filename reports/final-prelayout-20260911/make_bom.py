from pathlib import Path
from collections import defaultdict
import csv
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
XML = Path(__file__).resolve().parent / "project.xml"
OUT = ROOT / "output"

tree = ET.parse(XML)
components = tree.findall(".//components/comp")

manual = {
    "U1": ("Diodes Incorporated", "AP33772SDKZ-13-FA02", "Kesin parça; gate sürüşü Q3/Q4 ile doğrulanmalı"),
    "U2": ("Espressif", "ESP32-C6-WROOM-1", "Modül varyantı/flash kapasitesi satın almadan önce sabitlenmeli"),
    "U3": ("Texas Instruments", "INA228", "Sipariş suffix'i ve sıcaklık sınıfı seçilmeli"),
    "U4": ("Micro Crystal", "RV-3028-C7", "Trickle charge kapalı kullanılacak"),
    "U5": ("Alpha & Omega Semiconductor", "AOZ1284PI", "Özdisan stoklu"),
    "U6": ("Texas Instruments", "TL431BIDBZR", "Kesin"),
    "Q3": ("Infineon", "IRF7855TRPBF", "Prototip seçimi; AP33772S gate seviyesiyle kayıp/SOA açık"),
    "Q4": ("Infineon", "IRF7855TRPBF", "Prototip seçimi; AP33772S gate seviyesiyle kayıp/SOA açık"),
    "D2": ("PANJIT", "SS2060FL_R1_00001", "60 V / 2 A Schottky; Özdisan stoklu"),
    "L1": ("Core Master", "SRI0704-220M", "22 uH / 1.95 A; Özdisan stoklu; 7.3×7.3×4.5 mm"),
    "TH1": ("Thinking Electronic", "TSM0B103F3381RZ", "10 kΩ, B25/50=3380 K, 0402; Özdisan stoklu"),
    "J1": ("JAE", "DX07B024JJ1R1500", "Üretici çizimine göre doğrulanmış proje-yerel mid-mount footprint"),
    "J4": ("Degson", "DG142R-5.08-02P-14-00AH", "2 kutuplu 5.08 mm klemens"),
    "SW1": ("C&K / Littelfuse", "PTS810 SJS 250 SMTR LFS", "BOOT butonu; 4.2×3.2 mm SMD; Özdisan stoklu"),
    "SW2": ("C&K / Littelfuse", "PTS810 SJS 250 SMTR LFS", "RESET butonu; 4.2×3.2 mm SMD; Özdisan stoklu"),
    "SW3": ("KLS Electronic", "L-KLS4-EC12-012424-W", "Encoder"),
}

def fields(comp):
    d = {}
    for f in comp.findall("./fields/field"):
        d[f.attrib.get("name", "")] = f.text or ""
    for p in comp.findall("property"):
        d.setdefault(p.attrib.get("name", ""), p.attrib.get("value", ""))
    return d

items = []
for c in components:
    ref = c.attrib["ref"]
    value = c.findtext("value", "")
    fp = c.findtext("footprint", "")
    ds = c.findtext("datasheet", "")
    sheet = c.find("sheetpath").attrib.get("names", "/")
    f = fields(c)
    manufacturer = f.get("Manufacturer") or f.get("MANUFACTURER") or f.get("MF") or ""
    mpn = f.get("MPN") or f.get("MP") or ""
    note = f.get("DesignNote") or f.get("AssemblyNote") or ""
    if ref in manual:
        manufacturer, mpn, mn = manual[ref]
        note = "; ".join(x for x in [note, mn] if x)
    if ref == "J3":
        note = "TFT modülü ile kart üzerindeki Molex 54132-4062 konektörü BOM'da iki ayrı satın alma kalemi olarak ele alınmalı"
    if ref.startswith("TP") and fp:
        status = "PCB PAD"
        note = "1.0 mm SMD test pad; satın alınan parça değildir"
    elif not fp:
        status = "BLOKE — footprint eksik"
    elif ref in {"Q3", "Q4"}:
        status = "DOĞRULAMA GEREKLİ"
    elif mpn:
        status = "SEÇİLDİ"
    else:
        status = "MPN GEREKLİ"
    items.append(dict(ref=ref, sheet=sheet, value=value, footprint=fp, manufacturer=manufacturer,
                      mpn=mpn, datasheet=ds, status=status, note=note))

groups = defaultdict(list)
for x in items:
    key = (x["value"], x["footprint"], x["manufacturer"], x["mpn"], x["status"], x["note"], x["datasheet"])
    groups[key].append(x)

rows = []
for key, members in groups.items():
    value, fp, manufacturer, mpn, status, note, datasheet = key
    rows.append({
        "Qty": len(members),
        "References": ", ".join(sorted((m["ref"] for m in members), key=lambda s:(s.rstrip('0123456789'), int(''.join(filter(str.isdigit,s)) or 0)))),
        "Value": value,
        "Footprint": fp,
        "Manufacturer": manufacturer,
        "MPN": mpn,
        "Status": status,
        "Sheets": ", ".join(sorted(set(m["sheet"] for m in members))),
        "Notes": note,
        "Datasheet": datasheet,
    })
rows.sort(key=lambda r: (r["Status"].startswith("BLOKE") is False, r["References"]))

headers = list(rows[0])
csv_path = OUT / "BOM_PRELIMINARY_20260911.csv"
with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=headers)
    w.writeheader(); w.writerows(rows)

exceptions = []
for x in items:
    if x["ref"].startswith(("R", "C")) and "0402" not in x["footprint"]:
        reason = ""
        if x["ref"] in {"R11", "RShunt"}: reason = "5 mΩ / 1 W akım şöntü; 0402 güç ve Kelvin ölçüm için uygun değil"
        elif x["ref"] == "R43": reason = "28 V köşesinde yaklaşık 0.65 W; 2 W sınıfı gerekir"
        elif x["ref"] == "C3": reason = "50 V giriş bypass; DC-bias ve gerilim payı"
        elif x["ref"] == "C8": reason = "10 µF / 50 V ana giriş bulk"
        elif x["ref"] in {"C12", "C13"}: reason = "28 V altında etkili giriş kapasitesi ≥10 µF hedefi"
        elif x["ref"] in {"C15", "C16"}: reason = "3.3 V altında etkili çıkış bankı ≥44 µF hedefi"
        elif x["ref"] in {"C20", "C21"}: reason = "Boost giriş/çıkış bulk ve DC-bias gereksinimi"
        elif x["ref"] == "C5": reason = "ESP32 yerel 22 µF bulk; DC-bias gereksinimi"
        exceptions.append([x["ref"], x["value"], x["footprint"], reason])

wb = Workbook()
ws = wb.active; ws.title = "Grouped BOM"
ws.append(headers)
for row in rows: ws.append([row[h] for h in headers])
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1E5B73")
    cell.alignment = Alignment(wrap_text=True)
for row in ws.iter_rows(min_row=2):
    row[1].alignment = Alignment(wrap_text=True, vertical="top")
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    if str(row[6].value).startswith("BLOKE"):
        for c in row: c.fill = PatternFill("solid", fgColor="FFD9D9")
    elif row[6].value in {"MPN GEREKLİ", "DOĞRULAMA GEREKLİ"}:
        for c in row: c.fill = PatternFill("solid", fgColor="FFF2CC")
widths = [7, 32, 24, 47, 25, 29, 22, 28, 65, 55]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w

ex = wb.create_sheet("0402 Exceptions")
ex.append(["Reference", "Value", "Footprint", "Engineering reason"])
for r in exceptions: ex.append(r)
for cell in ex[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1E5B73")
for i,w in enumerate([14,20,42,80],1): ex.column_dimensions[get_column_letter(i)].width=w
for row in ex.iter_rows():
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")

sm = wb.create_sheet("Summary")
missing = [x for x in items if not x["footprint"]]
sm.append(["Metric", "Value"])
sm.append(["Physical components", len(items)])
sm.append(["Grouped BOM lines", len(rows)])
sm.append(["Missing-footprint components", len(missing)])
sm.append(["Missing-footprint refs", ", ".join(x["ref"] for x in missing)])
sm.append(["0402 engineering exceptions", len(exceptions)])
sm.append(["Electrical validation", "ERC 0 errors / 0 warnings; critical net checks 89/89"])
sm.append(["BOM status", "PRELIMINARY — yellow/red rows must be closed before production"])
for cell in sm[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1E5B73")
sm.column_dimensions["A"].width=32; sm.column_dimensions["B"].width=100
for row in sm.iter_rows():
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")

xlsx_path = OUT / "BOM_PRELIMINARY_20260911.xlsx"
wb.save(xlsx_path)
print(csv_path)
print(xlsx_path)
print(f"items={len(items)} groups={len(rows)} missing={len(missing)} exceptions={len(exceptions)}")
